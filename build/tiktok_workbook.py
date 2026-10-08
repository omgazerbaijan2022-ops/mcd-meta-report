"""Builds the TikTok analysis workbook with live formulas over the cleaned ad-level data.

Run after tiktok_analysis.py:  python3 tiktok_workbook.py
Then recalculate (LibreOffice) so cached values exist for viewers that do not compute formulas.
"""
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from tt_load import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT = os.environ.get("WB_OUT", os.path.join(HERE, "..", "tiktok", "McDonalds_AZ_TikTok_Analysis_Q2_vs_Q3_2026.xlsx"))
os.makedirs(os.path.dirname(OUT), exist_ok=True)

ARIAL = "Arial"
F_BASE = Font(name=ARIAL, size=10)
F_INPUT = Font(name=ARIAL, size=10, color="0000FF")          # hardcoded inputs (from the exports)
F_LINK = Font(name=ARIAL, size=10, color="008000")           # pulls from another sheet
F_HEAD = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=ARIAL, size=14, bold=True, color="27251F")
F_SUB = Font(name=ARIAL, size=10, italic=True, color="6E6E6E")
F_BOLD = Font(name=ARIAL, size=10, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="DA291C")
FILL_GOLD = PatternFill("solid", fgColor="FFBC0D")
FILL_KEY = PatternFill("solid", fgColor="FFF1C7")
FILL_ALT = PatternFill("solid", fgColor="F6F5F2")
THIN = Border(bottom=Side(style="thin", color="E0E0E0"))
USD2, USD3, USD0 = '$#,##0.00;($#,##0.00);"-"', '$#,##0.000;($#,##0.000);"-"', '$#,##0;($#,##0);"-"'
INT, PCT2, PCT1, DEC2 = '#,##0;(#,##0);"-"', '0.00%;-0.00%;"-"', '0.0%;-0.0%;"-"', '0.00;-0.00;"-"'

wb = Workbook()

def header(ws, row, labels, widths=None):
    for j, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w

def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = F_TITLE
    ws["A2"] = sub; ws["A2"].font = F_SUB

def put(ws, r, c, v, font=F_BASE, fmt=None, fill=None, bold=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = Font(name=ARIAL, size=10, bold=True, color=font.color) if bold else font
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    return cell

# ---------------------------------------------------------------- README
ws = wb.active; ws.title = "README"
title(ws, "McDonald's Azerbaijan – TikTok Ads: BEFORE vs AFTER analysis",
      "BEFORE = 1 Apr – 30 Jun 2026 (previous agency) · AFTER = 1 Jul – 30 Sep 2026 (current agency)")
notes = [
    ("Sources", "TikTok Ads Manager ad-level exports (USD): 'Tiktok_Ads_1_apr_-_30_June…20260401-20260630.xlsx' and 'Tiktok_Ads_1_july_-_Sep_30…20260701-20260930.xlsx'."),
    ("Active ads", f"Rows with zero spend (BEFORE {int((~B_all.active).sum())}, AFTER {int((~A_all.active).sum())}) have zero delivery; they are listed on 'Excluded rows' and left out of every calculation."),
    ("Weighted KPIs", "All rates are recalculated from summed totals with live formulas (CPM = spend ÷ impressions × 1,000; CTR = clicks ÷ impressions; CPC = spend ÷ clicks; view rate = 2-sec views ÷ impressions). Row-level rates are never averaged."),
    ("Unique reach", "TikTok's de-duplicated account reach comes from the 'Total of N results' row of each export (blue input on 'KPI Summary'). All other reach figures are ad-level reach summed, which counts a person once per ad."),
    ("Objective", "Inferred from campaign names (Reach, Community interaction, Video views, App promotion) and from Results = Reach. Opening_Nizami Mall reports results ≠ reach and is labelled 'Unlabelled (results ≠ reach)'."),
    ("Initiative", "Product keyword taken from the campaign name when it names a product, otherwise from the ad name (e.g. fifa, mixology, happy meal / HM, chicken / toyuq, cafe / cofe, grimace). Used only for like-for-like comparisons."),
    ("Results", "The 'Results' column mixes unnamed result types by objective, so it is shown per row but never summed across objectives."),
    ("Not in exports", "6-sec views, likes, comments, shares, engagement rate, conversions, daily breakdown and placement split are not available."),
    ("Colour code", "Blue text = values typed from the exports · black = formulas · green = links to another sheet · yellow fill = key KPI rows."),
    ("Sheets", "KPI Summary · By Objective · Campaigns BEFORE / AFTER · Ad Groups BEFORE / AFTER · Like-for-like · Data BEFORE / AFTER (ads, with per-ad KPIs) · Validation · Excluded rows."),
]
for i, (k, v) in enumerate(notes, 4):
    put(ws, i, 1, k, F_BOLD); c = put(ws, i, 2, v); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[i].height = 42
ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 120

# ---------------------------------------------------------------- Data sheets
DCOLS = ["Campaign", "Ad group", "Ad", "Objective", "Initiative", "Spend (USD)", "Reach (ad level)", "Impressions", "Clicks (destination)",
         "2-sec video views", "Results (as reported)", "Frequency (as reported)", "CPM (USD)", "CTR", "CPC (USD)", "2-sec view rate",
         "Cost / 1K 2-sec views (USD)", "Cost / 1K reached (USD)"]
DW = [34, 30, 40, 22, 18, 12, 14, 14, 12, 14, 14, 11, 10, 9, 10, 10, 12, 12]
DATA = {}
for p, df in [("BEFORE", B), ("AFTER", A)]:
    ws = wb.create_sheet(f"Data {p}")
    d = df.sort_values("spend", ascending=False).reset_index(drop=True)
    header(ws, 1, DCOLS, DW)
    for i, r in d.iterrows():
        row = i + 2
        vals = [r.campaign, r.adgroup, r.ad, r.objective, r.theme, r.spend, r.reach, r.impr, r.clicks, r.v2s, r.results, r.freq]
        for j, v in enumerate(vals, 1):
            put(ws, row, j, float(v) if isinstance(v, (np.floating,)) else (int(v) if isinstance(v, (np.integer,)) else v),
                F_INPUT, [None, None, None, None, None, USD2, INT, INT, INT, INT, INT, DEC2][j - 1])
        put(ws, row, 13, f'=IFERROR(F{row}/H{row}*1000,"")', fmt=USD3)
        put(ws, row, 14, f'=IFERROR(I{row}/H{row},"")', fmt=PCT2)
        put(ws, row, 15, f'=IF(I{row}>0,F{row}/I{row},"")', fmt=USD3)
        put(ws, row, 16, f'=IFERROR(J{row}/H{row},"")', fmt=PCT1)
        put(ws, row, 17, f'=IFERROR(F{row}/J{row}*1000,"")', fmt=USD2)
        put(ws, row, 18, f'=IFERROR(F{row}/G{row}*1000,"")', fmt=USD2)
    n = len(d) + 1
    ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:R{n}"
    DATA[p] = (f"'Data {p}'", n)

def rng(p, col):
    sh, n = DATA[p]
    return f"{sh}!${col}$2:${col}${n}"

# ---------------------------------------------------------------- KPI Summary
ws = wb.create_sheet("KPI Summary", 1)
title(ws, "KPI summary – BEFORE vs AFTER", "Totals over active ads; rates recalculated from totals. Blue = input from the export's Total row.")
header(ws, 4, ["KPI", "BEFORE (Apr–Jun)", "AFTER (Jul–Sep)", "Absolute change", "% change", "Better when", "Verdict"], [34, 18, 18, 18, 12, 14, 14])
rows = [
    ("Amount spent (USD)", f"=SUM({rng('BEFORE','F')})", f"=SUM({rng('AFTER','F')})", USD2, "Context"),
    ("Unique reach (de-duplicated)", float(B_tot.reach), float(A_tot.reach), INT, "Higher"),
    ("Impressions", f"=SUM({rng('BEFORE','H')})", f"=SUM({rng('AFTER','H')})", INT, "Context"),
    ("Frequency (impressions ÷ unique reach)", "=B7/B6", "=C7/C6", DEC2, "Context"),
    ("Cost per 1,000 unique reached (USD)", "=B5/B6*1000", "=C5/C6*1000", USD2, "Lower"),
    ("CPM (USD)", "=B5/B7*1000", "=C5/C7*1000", USD3, "Lower"),
    ("Clicks (destination)", f"=SUM({rng('BEFORE','I')})", f"=SUM({rng('AFTER','I')})", INT, "Higher"),
    ("CTR", "=B11/B7", "=C11/C7", PCT2, "Higher"),
    ("CPC (USD)", "=B5/B11", "=C5/C11", USD3, "Lower"),
    ("2-sec video views", f"=SUM({rng('BEFORE','J')})", f"=SUM({rng('AFTER','J')})", INT, "Higher"),
    ("2-sec view rate", "=B14/B7", "=C14/C7", PCT1, "Higher"),
    ("Cost per 1,000 2-sec views (USD)", "=B5/B14*1000", "=C5/C14*1000", USD2, "Lower"),
    ("Reach (sum of ad-level reach)", f"=SUM({rng('BEFORE','G')})", f"=SUM({rng('AFTER','G')})", INT, "Context"),
    ("Cost per 1,000 reached – ad level (USD)", "=B5/B17*1000", "=C5/C17*1000", USD2, "Lower"),
    ("Active ads", f"=COUNTA({rng('BEFORE','C')})", f"=COUNTA({rng('AFTER','C')})", INT, "Context"),
]
KEY = {"Amount spent (USD)", "Unique reach (de-duplicated)", "Cost per 1,000 unique reached (USD)", "CTR", "CPC (USD)", "2-sec view rate", "Cost per 1,000 2-sec views (USD)"}
for i, (lab, b, a, fmt, better) in enumerate(rows, 5):
    fill = FILL_KEY if lab in KEY else None
    put(ws, i, 1, lab, F_BOLD if fill else F_BASE, fill=fill)
    for col, v in ((2, b), (3, a)):
        is_input = not (isinstance(v, str) and v.startswith("="))
        font = F_INPUT if is_input else (F_LINK if "Data" in str(v) else F_BASE)
        c = put(ws, i, col, v, font, fmt, fill)
        if is_input:
            c.comment = Comment(f"Source: 'Total of N results' row (Reach column) of the {'BEFORE' if col == 2 else 'AFTER'} export – TikTok's de-duplicated reach.", "analysis")
    put(ws, i, 4, f"=C{i}-B{i}", fmt=fmt, fill=fill)
    put(ws, i, 5, f'=IFERROR((C{i}-B{i})/B{i},"")', fmt=PCT1, fill=fill)
    put(ws, i, 6, better, fill=fill)
    put(ws, i, 7, f'=IF(F{i}="Context","Context",IF(OR(AND(F{i}="Higher",E{i}>0),AND(F{i}="Lower",E{i}<0)),"Improvement","Decline"))', fill=fill)
ws.freeze_panes = "B5"
r0 = 5 + len(rows) + 1
put(ws, r0, 1, "Reading the table", F_BOLD)
for k, t in enumerate([
    "Spend fell ~10% while unique reach rose ~14%: cost per 1,000 unique people fell ~21% at an almost flat CPM – a genuine efficiency gain.",
    "Clicks (-17%) and 2-sec views (-30%) fell more than spend: CTR, CPC, view rate and cost per view all declined.",
    "Frequency is shown as context: lower means less repetition per person, not automatically better.",
], 1):
    put(ws, r0 + k, 1, t, F_SUB)

# ---------------------------------------------------------------- grouped sheets (SUMIFS)
GCOLS = ["Spend (USD)", "Share of spend", "Reach (ad level)", "Impressions", "Clicks", "2-sec views", "Frequency (ad level)", "CPM (USD)", "CTR", "CPC (USD)",
         "2-sec view rate", "Cost / 1K 2-sec views (USD)", "Cost / 1K reached (USD)"]
GFMT = [USD2, PCT1, INT, INT, INT, INT, DEC2, USD3, PCT2, USD3, PCT1, USD2, USD2]
SRC_COL = {"Spend (USD)": "F", "Reach (ad level)": "G", "Impressions": "H", "Clicks": "I", "2-sec views": "J"}

def grouped(name, p, keys, keyvals, extra_label_cols, sub):
    ws = wb.create_sheet(name)
    title(ws, name, sub)
    lab = list(keys) + extra_label_cols + ["Ads"] + GCOLS
    header(ws, 4, lab, [34] * len(keys) + [24] * len(extra_label_cols) + [7] + [13] * len(GCOLS))
    kc = {"campaign": "A", "adgroup": "B", "objective": "D", "theme": "E"}
    first = 5; last = first + len(keyvals) - 1
    for i, kv in enumerate(keyvals, first):
        vals = kv["keys"]
        for j, v in enumerate(vals, 1): put(ws, i, j, v, F_BOLD if j == 1 else F_BASE)
        for j, v in enumerate(kv.get("extra", []), len(vals) + 1): put(ws, i, j, v)
        crit = ",".join(f"{rng(p, kc[k])},${get_column_letter(j)}{i}" for j, k in enumerate(keys_internal[name], 1))
        base = len(vals) + len(extra_label_cols)
        put(ws, i, base + 1, f"=COUNTIFS({crit})", fmt=INT)
        col = {}
        for j, g in enumerate(GCOLS, base + 2):
            col[g] = get_column_letter(j)
        for g in GCOLS:
            j = ws[col[g] + "1"].column
            if g in SRC_COL:
                f = f"=SUMIFS({rng(p, SRC_COL[g])},{crit})"
            elif g == "Share of spend":
                f = f"={col['Spend (USD)']}{i}/SUM({rng(p, 'F')})"
            elif g == "Frequency (ad level)":
                f = f'=IFERROR({col["Impressions"]}{i}/{col["Reach (ad level)"]}{i},"")'
            elif g == "CPM (USD)":
                f = f'=IFERROR({col["Spend (USD)"]}{i}/{col["Impressions"]}{i}*1000,"")'
            elif g == "CTR":
                f = f'=IFERROR({col["Clicks"]}{i}/{col["Impressions"]}{i},"")'
            elif g == "CPC (USD)":
                f = f'=IF({col["Clicks"]}{i}>0,{col["Spend (USD)"]}{i}/{col["Clicks"]}{i},"")'
            elif g == "2-sec view rate":
                f = f'=IFERROR({col["2-sec views"]}{i}/{col["Impressions"]}{i},"")'
            elif g == "Cost / 1K 2-sec views (USD)":
                f = f'=IFERROR({col["Spend (USD)"]}{i}/{col["2-sec views"]}{i}*1000,"")'
            else:
                f = f'=IFERROR({col["Spend (USD)"]}{i}/{col["Reach (ad level)"]}{i}*1000,"")'
            put(ws, i, j, f, F_LINK if g in SRC_COL else F_BASE, GFMT[GCOLS.index(g)], FILL_ALT if (i - first) % 2 else None)
    # total row
    t = last + 1
    put(ws, t, 1, "Total", F_BOLD)
    base = len(keys) + len(extra_label_cols)
    put(ws, t, base + 1, f"=SUM({get_column_letter(base+1)}{first}:{get_column_letter(base+1)}{last})", fmt=INT, bold=True)
    for g in SRC_COL:
        c = get_column_letter(base + 2 + GCOLS.index(g))
        put(ws, t, base + 2 + GCOLS.index(g), f"=SUM({c}{first}:{c}{last})", fmt=GFMT[GCOLS.index(g)], bold=True)
    ws.freeze_panes = ws.cell(row=5, column=len(keys) + 1)
    ws.auto_filter.ref = f"A4:{get_column_letter(len(lab))}{last}"
    return ws

keys_internal = {}
for p, df in [("BEFORE", B), ("AFTER", A)]:
    nm = f"Campaigns {p}"; keys_internal[nm] = ["campaign"]
    cs = df.groupby("campaign").agg(spend=("spend", "sum"), obj=("objective", lambda s: "/".join(sorted(set(s))))).sort_values("spend", ascending=False)
    grouped(nm, p, ["Campaign"], [{"keys": [c], "extra": [r.obj]} for c, r in cs.iterrows()], ["Objective"],
            "SUMIFS over the ad-level data; rates recalculated from totals. Sorted by spend.")
for p, df in [("BEFORE", B), ("AFTER", A)]:
    nm = f"Ad Groups {p}"; keys_internal[nm] = ["campaign", "adgroup"]
    gs = df.groupby(["campaign", "adgroup"]).agg(spend=("spend", "sum"), obj=("objective", lambda s: "/".join(sorted(set(s))))).sort_values("spend", ascending=False)
    grouped(nm, p, ["Campaign", "Ad group"], [{"keys": list(k), "extra": [r.obj]} for k, r in gs.iterrows()], ["Objective"],
            "Ad group = (campaign, ad group) pair. BEFORE ad groups are one per campaign, named 'all'.")

# By objective: both periods side by side
ws = wb.create_sheet("By Objective", 2)
title(ws, "By objective – BEFORE vs AFTER", "Objective inferred from campaign names / results (see README). SUMIFS over the ad-level data.")
OBJS = ["Reach", "Video views", "Community interaction", "Unlabelled (results ≠ reach)", "App promotion"]
OC = ["Spend (USD)", "Share of spend", "Impressions", "Reach (ad level)", "Clicks", "2-sec views", "CPM (USD)", "CTR", "CPC (USD)", "2-sec view rate", "Cost / 1K 2-sec views (USD)", "Cost / 1K reached (USD)"]
OF = [USD2, PCT1, INT, INT, INT, INT, USD3, PCT2, USD3, PCT1, USD2, USD2]
header(ws, 4, ["Objective", "Period"] + OC, [30, 10] + [13] * len(OC))
r = 5
for o in OBJS:
    for p in ["BEFORE", "AFTER"]:
        put(ws, r, 1, o, F_BOLD); put(ws, r, 2, p)
        crit = f'{rng(p,"D")},$A{r}'
        put(ws, r, 3, f"=SUMIFS({rng(p,'F')},{crit})", F_LINK, USD2)
        put(ws, r, 4, f"=C{r}/SUM({rng(p,'F')})", fmt=PCT1)
        put(ws, r, 5, f"=SUMIFS({rng(p,'H')},{crit})", F_LINK, INT)
        put(ws, r, 6, f"=SUMIFS({rng(p,'G')},{crit})", F_LINK, INT)
        put(ws, r, 7, f"=SUMIFS({rng(p,'I')},{crit})", F_LINK, INT)
        put(ws, r, 8, f"=SUMIFS({rng(p,'J')},{crit})", F_LINK, INT)
        put(ws, r, 9, f'=IFERROR(C{r}/E{r}*1000,"")', fmt=USD3)
        put(ws, r, 10, f'=IFERROR(G{r}/E{r},"")', fmt=PCT2)
        put(ws, r, 11, f'=IF(G{r}>0,C{r}/G{r},"")', fmt=USD3)
        put(ws, r, 12, f'=IFERROR(H{r}/E{r},"")', fmt=PCT1)
        put(ws, r, 13, f'=IFERROR(C{r}/H{r}*1000,"")', fmt=USD2)
        put(ws, r, 14, f'=IFERROR(C{r}/F{r}*1000,"")', fmt=USD2)
        if p == "AFTER":
            for c in range(1, 15): ws.cell(row=r, column=c).border = THIN
        r += 1
ws.freeze_panes = "C5"

# Like-for-like
ws = wb.create_sheet("Like-for-like", 3)
title(ws, "Like-for-like initiatives – Reach-objective ads only",
      "Matched only where the product keyword appears in both periods. Mixology and Happy Meal rest on one BEFORE ad each.")
LC = ["Initiative", "BEFORE ads", "AFTER ads", "BEFORE spend", "AFTER spend", "BEFORE cost/1K reached", "AFTER cost/1K reached", "% change cost/1K",
      "BEFORE CPM", "AFTER CPM", "BEFORE CTR", "AFTER CTR", "% change CTR", "BEFORE view rate", "AFTER view rate", "% change view rate",
      "BEFORE freq (ad level)", "AFTER freq (ad level)"]
header(ws, 4, LC, [20] + [11] * (len(LC) - 1))
for i, t in enumerate(["FIFA / World Cup", "Mixology", "Grimace", "Happy Meal", "Chicken", "McCafé"], 5):
    put(ws, i, 1, t, F_BOLD)
    cb = f'{rng("BEFORE","E")},$A{i},{rng("BEFORE","D")},"Reach"'; ca = f'{rng("AFTER","E")},$A{i},{rng("AFTER","D")},"Reach"'
    S = lambda p, col, c: f"SUMIFS({rng(p,col)},{c})"
    put(ws, i, 2, f"=COUNTIFS({cb})", fmt=INT); put(ws, i, 3, f"=COUNTIFS({ca})", fmt=INT)
    put(ws, i, 4, "=" + S("BEFORE", "F", cb), F_LINK, USD2); put(ws, i, 5, "=" + S("AFTER", "F", ca), F_LINK, USD2)
    put(ws, i, 6, f"=D{i}/{S('BEFORE','G',cb)}*1000", fmt=USD3); put(ws, i, 7, f"=E{i}/{S('AFTER','G',ca)}*1000", fmt=USD3)
    put(ws, i, 8, f"=(G{i}-F{i})/F{i}", fmt=PCT1)
    put(ws, i, 9, f"=D{i}/{S('BEFORE','H',cb)}*1000", fmt=USD3); put(ws, i, 10, f"=E{i}/{S('AFTER','H',ca)}*1000", fmt=USD3)
    put(ws, i, 11, f"={S('BEFORE','I',cb)}/{S('BEFORE','H',cb)}", fmt=PCT2); put(ws, i, 12, f"={S('AFTER','I',ca)}/{S('AFTER','H',ca)}", fmt=PCT2)
    put(ws, i, 13, f"=(L{i}-K{i})/K{i}", fmt=PCT1)
    put(ws, i, 14, f"={S('BEFORE','J',cb)}/{S('BEFORE','H',cb)}", fmt=PCT1); put(ws, i, 15, f"={S('AFTER','J',ca)}/{S('AFTER','H',ca)}", fmt=PCT1)
    put(ws, i, 16, f"=(O{i}-N{i})/N{i}", fmt=PCT1)
    put(ws, i, 17, f"={S('BEFORE','H',cb)}/{S('BEFORE','G',cb)}", fmt=DEC2); put(ws, i, 18, f"={S('AFTER','H',ca)}/{S('AFTER','G',ca)}", fmt=DEC2)
ws.freeze_panes = "B5"

# Validation
ws = wb.create_sheet("Validation")
title(ws, "Validation – ad rows vs the export's Total row", "Blue = typed from each export's 'Total of N results' row. Difference should be 0.")
header(ws, 4, ["Metric", "Period", "Total row (export)", "Sum of ad rows", "Difference"], [22, 10, 20, 20, 14])
r = 5
for p, tot in [("BEFORE", B_tot), ("AFTER", A_tot)]:
    for lab, col, val, fmt in [("Spend (USD)", "F", tot.spend, USD2), ("Impressions", "H", tot.impr, INT), ("Clicks", "I", tot.clicks, INT), ("2-sec video views", "J", tot.v2s, INT)]:
        put(ws, r, 1, lab); put(ws, r, 2, p); put(ws, r, 3, float(val), F_INPUT, fmt)
        put(ws, r, 4, f"=SUM({rng(p,col)})", F_LINK, fmt); put(ws, r, 5, f"=ROUND(C{r}-D{r},2)", fmt=fmt)
        r += 1
put(ws, r + 1, 1, "Zero-spend rows excluded: these have zero impressions, clicks and views, so they do not affect the sums.", F_SUB)

# Excluded rows
ws = wb.create_sheet("Excluded rows")
title(ws, "Excluded rows – zero spend, zero delivery", "Legacy ads listed in the exports without any delivery in the period.")
header(ws, 4, ["Period", "Campaign", "Ad group", "Ad"], [10, 40, 34, 40])
r = 5
for p, al in [("BEFORE", B_all), ("AFTER", A_all)]:
    for _, x in al[~al.active].iterrows():
        put(ws, r, 1, p); put(ws, r, 2, x.campaign, F_INPUT); put(ws, r, 3, x.adgroup, F_INPUT); put(ws, r, 4, x.ad, F_INPUT); r += 1

order = ["README", "KPI Summary", "By Objective", "Like-for-like", "Campaigns BEFORE", "Campaigns AFTER", "Ad Groups BEFORE", "Ad Groups AFTER", "Data BEFORE", "Data AFTER", "Validation", "Excluded rows"]
wb._sheets = [wb[n] for n in order]
for w in wb.worksheets:
    w.sheet_view.showGridLines = False
wb["KPI Summary"].sheet_properties.tabColor = "DA291C"
wb.active = 1
wb.save(OUT)
print("wrote", OUT)
